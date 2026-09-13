"""Local OSRM integration check with invented nearby pickup points, no API keys."""
import io
import json
import os
from pathlib import Path
import sys
import tempfile

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))


def main():
    with tempfile.TemporaryDirectory(prefix='daonroad-road-check-') as cache:
        os.environ['DAONROAD_CACHE_DIR']=cache
        os.environ['OSRM_BASE_URL']='http://127.0.0.1:5001'
        from api.models import RouteRequest
        from api.routing import calculate_route
        from export.excel_exporter import ExcelExporter
        from openpyxl import load_workbook
        request=RouteRequest(
            passengers=[dict(name=f'Synthetic {i}',address=f'Test {i}',passenger_count=1,
                             lat=37.5+(i%4)*.002,lng=127.+(i//4)*.002) for i in range(12)],
            vehicles=[dict(bus_id=f'Bus {i}',capacity=21,start_location='Test depot',
                           start_lat=37.51,start_lng=127.01) for i in range(2)],
            destination='Test school',destination_lat=37.498,destination_lng=127.03,
            arrival_time='10:00',service_date='2026-09-13',options={'search_seconds':2})
        result=calculate_route(request)
        assert result['matrix_source']=='osrm'
        assert result['summary']['total_passengers']==12
        workbook=load_workbook(io.BytesIO(ExcelExporter().export(result['routes'],result['destination'],result['summary'])))
        assert workbook.sheetnames==['Bus Summary','Route Detail','Passenger']
        report={'source':result['matrix_source'],'summary':result['summary'],
                'warnings':result['warnings'],'excel_sheets':workbook.sheetnames}
        text=json.dumps(report,ensure_ascii=False,indent=2)
        print(text)
        (Path(__file__).resolve().parents[1]/'docs'/'road-integration.json').write_text(text+'\n',encoding='utf-8')


if __name__=='__main__':
    main()
